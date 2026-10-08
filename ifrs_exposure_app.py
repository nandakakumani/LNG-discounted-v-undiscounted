from io import BytesIO
import re
import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.utils import get_column_letter

YEARS = list(range(2026, 2033))
DIVISORS = {"NET GAS (TBTU)": 1_000_000.0, "NET OIL (MM BBL)": 1_000_000.0, "HH": 1_000_000.0}
REQUIRED = {"PnL Year", "Delta", "Final Index", "Group", "ViewGroup"}
OIL = {"BRENT", "BRENT BULLET", "BRENT FUTURES", "BRENT SWAPS", "DATED BRENT", "DUBAI", "JCC", "JCC SWAPS", "WTI"}
GAS = {"TTF", "NBP", "PEG", "PEG - DNK", "PVB", "THE", "TVB", "ZTP", "ZTP - DNK", "JKM"}


def clean(v):
    return "" if pd.isna(v) else re.sub(r"\s+", " ", str(v).strip()).upper()


def parse_year(s):
    out = pd.Series(pd.NA, index=s.index, dtype="Int64")
    n = pd.to_numeric(s, errors="coerce")
    direct = n.between(1900, 2200)
    out.loc[direct] = n.loc[direct].round().astype("Int64")
    serial = n.between(20000, 100000) & ~direct
    if serial.any():
        out.loc[serial] = pd.to_datetime(n.loc[serial], unit="D", origin="1899-12-30", errors="coerce").dt.year.astype("Int64")
    remaining = out.isna()
    if remaining.any():
        out.loc[remaining] = pd.to_datetime(s.loc[remaining], errors="coerce").dt.year.astype("Int64")
    return out


def category(row):
    group, index = clean(row.get("Group")), clean(row.get("Final Index"))
    if index == "HH" or index.startswith("HENRY HUB"):
        return "HH"
    if "EUA" in group or "EUA" in index or "FREIGHT" in group or "FREIGHT" in index:
        return "EXCLUDED"
    if group == "OIL" or index in OIL or any(x in index for x in ("BRENT", "JCC", "DUBAI")):
        return "NET OIL (MM BBL)"
    gas_group = group in {"NATURAL GAS & LNG", "UNSOLD HOUSE CURVES"} or "NATURAL GAS" in group or "LNG" in group or "GAS" in group
    if gas_group or index in GAS or "LNG" in index:
        return "NET GAS (TBTU)"
    return "UNCLASSIFIED"


def load_data(file):
    xls = pd.ExcelFile(file, engine="openpyxl")
    matches = [s for s in xls.sheet_names if s.strip().casefold() == "raw data"]
    if not matches:
        raise ValueError('Workbook must contain a sheet named "Raw Data".')
    df = pd.read_excel(xls, sheet_name=matches[0])
    df.columns = [str(c).strip() for c in df.columns]
    missing = REQUIRED - set(df.columns)
    if missing:
        raise ValueError("Missing columns: " + ", ".join(sorted(missing)))
    df["Delta"] = pd.to_numeric(df["Delta"], errors="coerce").fillna(0.0)
    df["IFRS Year"] = parse_year(df["PnL Year"])
    df["Category"] = df.apply(category, axis=1)
    df["View"] = df["ViewGroup"].map(clean)
    return df


def make_table(df, view, years):
    d = df[df["View"].isin([view.upper(), "BOTH"])]
    rows = []
    for cat in ["NET GAS (TBTU)", "NET OIL (MM BBL)", "HH"]:
        values = [d.loc[(d["Category"] == cat) & (d["IFRS Year"] == y), "Delta"].sum() / DIVISORS[cat] for y in years]
        rows.append([cat] + values + [sum(values)])
    return pd.DataFrame(rows, columns=["LNG PORTFOLIO EXPOSURES"] + [f"IFRS {y}" for y in years] + ["NET"]).set_index("LNG PORTFOLIO EXPOSURES")


def write_section(ws, title_row, title, table, years, purple_word=None, impact_refs=None):
    black = "000000"; purple = "7030A0"; red = "FF0000"; grey = "F2F2F2"
    thin = Side(style="thin", color=black)
    all_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Exact layout from the supplied template: title C:L, table C:J, spacer K, NET L.
    ws.merge_cells(start_row=title_row, start_column=3, end_row=title_row, end_column=12)
    tc = ws.cell(title_row, 3)
    if title.endswith("(UNDISCOUNTED)") or title.endswith("(DISCOUNTED)"):
        prefix, suffix = title.rsplit(" ", 1)
        tc.value = CellRichText(
            TextBlock(InlineFont(rFont="Arial", sz=12, b=True, color=black), prefix + " "),
            TextBlock(InlineFont(rFont="Arial", sz=12, b=True, color=purple), suffix),
        )
    else:
        tc.value = title
    tc.font = Font(name="Arial", size=12, bold=True, color=black)
    tc.alignment = Alignment(horizontal="center", vertical="center")
    tc.fill = PatternFill("solid", fgColor=grey)
    for col in range(3, 13):
        ws.cell(title_row, col).border = Border(top=thin, bottom=thin)
    ws.cell(title_row, 3).border = Border(left=thin, top=thin, bottom=thin)
    ws.cell(title_row, 12).border = Border(right=thin, top=thin, bottom=thin)

    header_row = title_row + 2
    headers = ["LNG PORTFOLIO EXPOSURES"] + [f"IFRS {y}" for y in years]
    for col, val in zip(range(3, 11), headers):
        c = ws.cell(header_row, col, val)
        c.font = Font(name="Arial", size=11, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = all_border
    net_h = ws.cell(header_row, 12, "NET")
    net_h.font = Font(name="Arial", size=11, bold=True)
    net_h.alignment = Alignment(horizontal="center", vertical="center")
    net_h.border = all_border

    start = header_row + 1
    row_map = {}
    for i, label in enumerate(table.index):
        r = start + i; row_map[label] = r
        lab = ws.cell(r, 3, label)
        lab.font = Font(name="Arial", size=11, bold=True)
        lab.alignment = Alignment(horizontal="center")
        lab.border = all_border
        for idx, col in enumerate(range(4, 11)):
            c = ws.cell(r, col)
            if impact_refs:
                dr, ur = impact_refs[label]
                c.value = f"={get_column_letter(col)}{dr}-{get_column_letter(col)}{ur}"
            else:
                c.value = float(table.iloc[i, idx])
            c.font = Font(name="Arial", size=11, color=red if (not impact_refs and c.value < 0) else black)
            c.alignment = Alignment(horizontal="center")
            c.number_format = '0.00;[Red]-0.00;0.00'
            c.border = all_border
        # K is intentionally blank as shown in the requested format.
        net = ws.cell(r, 12)
        if impact_refs:
            net.value = f"=SUM(D{r}:J{r})"
        else:
            net.value = float(table.iloc[i, -1])
        net.font = Font(name="Arial", size=11, bold=True, color=red if (not impact_refs and net.value < 0) else black)
        net.alignment = Alignment(horizontal="center")
        net.number_format = '0.00;[Red]-0.00;0.00'
        net.border = all_border
    return row_map


def excel_output(df, undisc, disc, impact, years):
    wb = Workbook(); ws = wb.active; ws.title = "Summary"; ws.sheet_view.showGridLines = True
    ws.column_dimensions["B"].width = 2
    ws.column_dimensions["C"].width = 34
    for col in "DEFGHIJ": ws.column_dimensions[col].width = 12.5
    ws.column_dimensions["K"].width = 8
    ws.column_dimensions["L"].width = 12
    for r in range(1, 26): ws.row_dimensions[r].height = 21

    ur = write_section(ws, 2, "OVERALL PORTFOLIO BALANCE BASE CASE (UNDISCOUNTED)", undisc, years)
    dr = write_section(ws, 9, "OVERALL PORTFOLIO BALANCE BASE CASE (DISCOUNTED)", disc, years)
    refs = {k: (dr[k], ur[k]) for k in ur}
    write_section(ws, 16, "DISCOUNTING IMPACT ON HEDGES", impact, years, impact_refs=refs)

    # Match the yellow review highlight visible in the template: HH / IFRS 2028 impact.
    ws["F21"].fill = PatternFill("solid", fgColor="FFFF00")

    def add_data_sheet(name, frame):
        sh = wb.create_sheet(name)
        sh.sheet_view.showGridLines = False
        export = frame.copy()
        for c, col_name in enumerate(export.columns, 1):
            cell = sh.cell(1, c, str(col_name))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
        for r, vals in enumerate(export.itertuples(index=False, name=None), 2):
            for c, val in enumerate(vals, 1):
                if isinstance(val, np.generic): val = val.item()
                if pd.isna(val): val = None
                sh.cell(r, c, val)
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = sh.dimensions
        for c, col_name in enumerate(export.columns, 1):
            width = max(12, len(str(col_name)) + 2)
            for val in export.iloc[:500, c - 1]:
                width = max(width, min(35, len(str(val)) + 2))
            sh.column_dimensions[get_column_letter(c)].width = width
        return sh

    # Product Detail keeps the auditable product/index breakdown used in the earlier version.
    detail_parts = []
    included = df[df["Category"].isin(DIVISORS)].copy()
    for view_name in ["Undiscounted", "Discounted"]:
        vd = included[included["View"].isin([view_name.upper(), "BOTH"])].copy()
        grouped = (vd[vd["IFRS Year"].isin(years)]
            .groupby(["Category", "Final Index", "IFRS Year"], dropna=False)["Delta"]
            .sum().reset_index())
        grouped["Exposure"] = grouped.apply(lambda r: r["Delta"] / DIVISORS[r["Category"]], axis=1)
        grouped.insert(0, "View", view_name)
        detail_parts.append(grouped.drop(columns="Delta"))
    product_detail = pd.concat(detail_parts, ignore_index=True)
    add_data_sheet("Product Detail", product_detail)

    audit = df.groupby(["Category", "Group", "Final Index"], dropna=False).agg(Rows=("Delta", "size"), Raw_Delta=("Delta", "sum")).reset_index()
    add_data_sheet("Classification Audit", audit)

    # Preserve all uploaded raw rows plus the derived IFRS year and classification fields.
    add_data_sheet("Filtered Raw Data", df)

    wb.calculation.fullCalcOnLoad = True; wb.calculation.forceFullCalc = True; wb.calculation.calcMode = "auto"
    out = BytesIO(); wb.save(out); return out.getvalue()


def app():
    st.set_page_config(page_title="LNG IFRS Exposure Builder", layout="wide")
    st.title("LNG IFRS Exposure Builder")
    uploaded = st.file_uploader('Upload workbook containing the "Raw Data" tab', type=["xlsx"])
    if not uploaded: return
    try: df = load_data(uploaded)
    except Exception as e: st.error(str(e)); return
    undisc = make_table(df, "Undiscounted", YEARS)
    disc = make_table(df, "Discounted", YEARS)
    impact = disc - undisc
    for title, table in [("Overall Portfolio Balance Base Case (Undiscounted)", undisc), ("Overall Portfolio Balance Base Case (Discounted)", disc), ("Discounting Impact on Hedges", impact)]:
        st.subheader(title); st.dataframe(table.style.format("{:.2f}"), use_container_width=True)
    output = excel_output(df, undisc, disc, impact, YEARS)
    st.download_button("Download formatted Excel output", output, "IFRS_Exposure_Output.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")

if __name__ == "__main__": app()
